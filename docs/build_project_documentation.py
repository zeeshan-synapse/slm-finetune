from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


PROJECT_DIR = Path(__file__).resolve().parents[1]
OUTPUT_PATH = PROJECT_DIR / "docs" / "Synapse_SLM_Project_Documentation.docx"

BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
MUTED = RGBColor(90, 98, 108)
LIGHT_FILL = "F2F4F7"
CALLOUT_FILL = "E8EEF5"


def set_run_font(run, size=11, bold=False, color=None, italic=False):
    run.font.name = "Calibri"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Calibri")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Calibri")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = color


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, width_dxa):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        grid.append(grid_col)

    for row in table.rows:
        for index, cell in enumerate(row.cells):
            set_cell_width(cell, widths[index])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    tr_pr.append(repeat)


def add_heading(doc, text, level=1):
    paragraph = doc.add_paragraph(style=f"Heading {level}")
    paragraph.add_run(text)
    return paragraph


def add_body(doc, text, bold_lead=None):
    paragraph = doc.add_paragraph()
    if bold_lead and text.startswith(bold_lead):
        lead = paragraph.add_run(bold_lead)
        set_run_font(lead, bold=True)
        text = text[len(bold_lead):]
    run = paragraph.add_run(text)
    set_run_font(run)
    return paragraph


def add_bullet(doc, text):
    paragraph = doc.add_paragraph(style="List Bullet")
    run = paragraph.add_run(text)
    set_run_font(run)
    return paragraph


def add_callout(doc, label, text):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [9360])
    cell = table.cell(0, 0)
    shade_cell(cell, CALLOUT_FILL)
    paragraph = cell.paragraphs[0]
    lead = paragraph.add_run(f"{label}: ")
    set_run_font(lead, bold=True, color=DARK_BLUE)
    run = paragraph.add_run(text)
    set_run_font(run)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_key_value_table(doc, rows):
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    set_table_geometry(table, [2600, 6760])
    header = table.rows[0].cells
    for cell, text in zip(header, ("Area", "Current implementation")):
        shade_cell(cell, LIGHT_FILL)
        run = cell.paragraphs[0].add_run(text)
        set_run_font(run, bold=True, color=DARK_BLUE)
    set_repeat_table_header(table.rows[0])

    for left, right in rows:
        cells = table.add_row().cells
        for cell, text in zip(cells, (left, right)):
            run = cell.paragraphs[0].add_run(text)
            set_run_font(run, size=10.5)
    set_table_geometry(table, [2600, 6760])
    return table


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    set_run_font(run, size=9, color=MUTED)
    field_begin = OxmlElement("w:fldChar")
    field_begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = "PAGE"
    field_end = OxmlElement("w:fldChar")
    field_end.set(qn("w:fldCharType"), "end")
    run._r.append(field_begin)
    run._r.append(instruction)
    run._r.append(field_end)


def configure_document(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1

    heading_tokens = {
        "Heading 1": (16, BLUE, 16, 8),
        "Heading 2": (13, BLUE, 12, 6),
        "Heading 3": (12, DARK_BLUE, 8, 4),
    }
    for name, (size, color, before, after) in heading_tokens.items():
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    bullet = styles["List Bullet"]
    bullet.font.name = "Calibri"
    bullet.font.size = Pt(11)
    bullet.paragraph_format.left_indent = Inches(0.5)
    bullet.paragraph_format.first_line_indent = Inches(-0.25)
    bullet.paragraph_format.space_after = Pt(5)
    bullet.paragraph_format.line_spacing = 1.1

    header = section.header.paragraphs[0]
    header.text = ""
    run = header.add_run("SYNAPSE SLM PROJECT | TECHNICAL OVERVIEW")
    set_run_font(run, size=9, bold=True, color=MUTED)
    footer = section.footer.paragraphs[0]
    add_page_number(footer)


def build_document():
    doc = Document()
    configure_document(doc)

    title = doc.add_paragraph()
    title.paragraph_format.space_before = Pt(8)
    title.paragraph_format.space_after = Pt(4)
    run = title.add_run("Synapse SLM Project Documentation")
    set_run_font(run, size=24, bold=True, color=DARK_BLUE)

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(16)
    run = subtitle.add_run("Local fine-tuning, FAISS RAG, model comparison, and grounded website assistance")
    set_run_font(run, size=12, italic=True, color=MUTED)

    add_callout(
        doc,
        "Current position",
        "The strongest working path is an instruct model combined with the local Synapse knowledge base. "
        "The project also contains two experimental LoRA fine-tunes for comparison and diagnosis.",
    )

    add_heading(doc, "1. Project Purpose")
    add_body(
        doc,
        "This project builds a fully local assistant for Synapse Tech website and domain questions. "
        "It prepares conversational data, fine-tunes a compact language model with LoRA, exports models to GGUF, "
        "runs inference through Ollama, and grounds answers using a local FAISS knowledge base."
    )
    add_body(
        doc,
        "The interactive development interface compares fine-tuned models, untrained base models with RAG, "
        "and plain base-model responses. This makes it possible to separate model-training problems from retrieval "
        "or prompting problems."
    )

    add_heading(doc, "2. Current Architecture")
    add_key_value_table(
        doc,
        [
            ("Base used for fine-tuning", "Qwen2.5-1.5B-Instruct from Hugging Face/MLX format."),
            ("Fine-tuned models", "synapse-1.5b-v1 and synapse-1.5b-v2, both derived from the 1.5B Qwen base."),
            ("Runtime comparison models", "Qwen2.5 1.5B, 3B, 7B, and llama3:latest are selectable in chat/chat.py."),
            ("Knowledge base", "Cleaned Synapse website content indexed with nomic-embed-text and stored in FAISS."),
            ("Inference", "Local Ollama API; GGUF models can also be used with llama.cpp for deployment."),
            ("Observability", "Debug metadata records selected model/profile, intent, rewrite query, sources, token use, and route."),
        ],
    )

    add_heading(doc, "Runtime Flow", level=2)
    add_body(
        doc,
        "For Base + RAG, the selected runtime model is used for semantic query handling and grounded answer "
        "generation. The embedding model remains separate because the FAISS index was built with a dedicated "
        "embedding model."
    )
    for step in (
        "User question enters chat/chat.py.",
        "The selected model produces a KB-oriented query and semantic intent metadata.",
        "The rewritten query is embedded with nomic-embed-text and searched against the FAISS index.",
        "Relevant chunks are reranked and selected as evidence.",
        "The selected model receives the original question plus evidence and generates the final answer.",
        "Guardrails, validation, and structured debug logging inspect the result.",
    ):
        add_bullet(doc, step)

    add_heading(doc, "3. Main Components")
    add_key_value_table(
        doc,
        [
            ("dataset/", "Prepares website Q&A, conversational samples, augmentation, and mixed training datasets."),
            ("train/finetune.py", "Formats training samples and runs MLX-LM LoRA training."),
            ("export/", "Fuses LoRA adapters, converts the model to GGUF, and quantizes it for local inference."),
            ("scripts/query_kb.py", "Builds query profiles, embeds search text, searches FAISS, and reranks evidence."),
            ("scripts/answer_with_kb.py", "Controls semantic query handling, retrieval, model profiles, answer generation, safety, and logging."),
            ("chat/kb_answer.py", "Bridge between the chat UI and the reusable KB/RAG pipeline."),
            ("chat/chat.py", "Development comparison UI with model, RAG behavior, and answer-path selection."),
            ("chat/assistant.py", "Cleaner single-assistant path intended as a production-facing foundation."),
            ("data/runtime/", "Editable product cards and policy wording kept outside Python logic."),
        ],
    )

    add_heading(doc, "4. Data and Model Artifacts")
    add_body(
        doc,
        "Website knowledge and model-training data are related but serve different purposes. Training datasets "
        "teach behavior and response patterns; the FAISS knowledge base supplies current Synapse facts at runtime."
    )
    add_bullet(doc, "Training data: data/mixed_dataset.jsonl and data/mixed_dataset_v2.jsonl.")
    add_bullet(doc, "Conversational data: data/ultrachat_sample.jsonl.")
    add_bullet(doc, "Knowledge-base sources and chunks: data/knowledge-base/source_docs.jsonl and chunks.jsonl.")
    add_bullet(doc, "Vector index: data/knowledge-base/faiss.index with index_meta.jsonl and index_manifest.json.")
    add_bullet(doc, "Fine-tuned adapters and merged models: models/qwen1.5b-finetuned*, models/qwen1.5b-merged*.")
    add_bullet(doc, "Deployment model: quantized GGUF files under models/gguf/.")

    add_heading(doc, "5. Running the Project")
    add_body(doc, "Start Ollama if it is not already running:")
    code = doc.add_paragraph()
    code.style = doc.styles["Normal"]
    code.paragraph_format.left_indent = Inches(0.25)
    shade = OxmlElement("w:shd")
    shade.set(qn("w:fill"), "F7F7F7")
    code._p.get_or_add_pPr().append(shade)
    run = code.add_run("ollama serve")
    set_run_font(run, size=10, color=DARK_BLUE)

    add_body(doc, "From the project root:")
    code = doc.add_paragraph()
    code.paragraph_format.left_indent = Inches(0.25)
    shade = OxmlElement("w:shd")
    shade.set(qn("w:fill"), "F7F7F7")
    code._p.get_or_add_pPr().append(shade)
    run = code.add_run("source venv/bin/activate\npython chat/chat.py")
    set_run_font(run, size=10, color=DARK_BLUE)

    add_body(
        doc,
        "At startup, select the pipeline model, deterministic or model-generated RAG behavior, and the answer path. "
        "Use batch mode to paste evaluation questions and END to execute them. Enable debug mode when inspecting "
        "intent, query rewrite, retrieved sources, token use, or latency."
    )

    add_heading(doc, "6. Current Findings")
    add_bullet(doc, "Base + RAG currently gives the best balance of grounded knowledge and natural conversation.")
    add_bullet(doc, "Qwen2.5 1.5B is fastest; Llama 3 often gives stronger language and recommendation quality.")
    add_bullet(doc, "Larger models are slower because each RAG question can involve semantic query handling and final generation.")
    add_bullet(doc, "Model-specific profiles now tune prompts, retrieval depth, temperature, and output budget for 3B, 7B, and Llama 3.")
    add_bullet(doc, "Product-catalog and recommendation routing were corrected so plural catalog questions are not treated as one-product requests.")
    add_bullet(doc, "The same FAISS evidence can produce different results depending on model classification and generation behavior.")

    doc.add_page_break()
    add_heading(doc, "7. Known Limitations and Risks")
    add_callout(
        doc,
        "Important training issue",
        "The current training formatter appends an extra assistant-start marker after completed conversations. "
        "Before another serious fine-tune, remove that marker, regenerate train/validation files, inspect samples, "
        "and retrain from a fresh adapter rather than resuming an old checkpoint.",
    )
    add_bullet(doc, "The website source includes marketing language and unsupported-looking percentages that can be repeated by RAG.")
    add_bullet(doc, "Intent and retrieval behavior still need evaluation on broader paraphrases and adversarial prompts.")
    add_bullet(doc, "Deterministic paths are reliable but can sound robotic; generated paths are natural but need stronger validation.")
    add_bullet(doc, "The current chat comparison UI is a development tool, not the final production interface.")
    add_bullet(doc, "Authentication, user permissions, tool governance, rate limits, and production monitoring are not yet complete.")

    add_heading(doc, "8. Recommended Next Steps")
    for step in (
        "Merge query rewriting and intent classification into one structured model call to reduce latency.",
        "Fix the training formatter and run a small clean LoRA experiment before full retraining.",
        "Create a larger automated evaluation set covering paraphrases, unsupported claims, recommendations, and prompt attacks.",
        "Clean or curate noisy website chunks and add reranking metrics such as recall@k and answer faithfulness.",
        "Keep Base + RAG as the demo baseline and only adopt a fine-tune when it beats that baseline consistently.",
        "For paid deployments, add authentication, permissions, audit logs, schema validation, and model/tool governance.",
    ):
        add_bullet(doc, step)

    add_heading(doc, "9. Practical Project Position")
    add_body(
        doc,
        "The project is a strong prototype for private RAG assistants and local AI deployments. Its reusable value "
        "is the end-to-end pipeline: data preparation, LoRA experimentation, model export, local inference, semantic "
        "retrieval, grounded generation, guardrails, model comparison, and observability. A larger model can improve "
        "quality, but production success will depend equally on data quality, retrieval, evaluation, security, and deployment."
    )

    doc.core_properties.title = "Synapse SLM Project Documentation"
    doc.core_properties.subject = "Current technical overview and project handoff"
    doc.core_properties.author = "Synapse SLM Project"
    doc.save(OUTPUT_PATH)
    print(OUTPUT_PATH)


if __name__ == "__main__":
    build_document()
